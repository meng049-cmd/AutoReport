import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import autoreport as a
FIXTURE_SHEETS = ['簡體對照', '工單連結', '合併欄位', '產品系列', '無對照']


def fixture():
    book = Workbook()
    s = book.active
    s.title = '簡體對照'
    s.append(['机种', '机型', '料号', '工单号'])
    s.append(['繁體測試機種', 'Goodix', '98000001-107', 'ORDER001'])
    s.append(['另一機種', 'FPC', '98000002-107', 'ORDER002'])
    s.append(['繁體測試機種', 'Synaptics', 'SHARED', 'ORDER003'])
    s.append(['另一機種', 'FPC', 'SHARED', 'ORDER004'])
    s.append(['繁體測試機種', 'Goodix', '98000001-108', 'ORDER005'])
    s.append(['不完整', 'N/A', 'UNKNOWN-IC', 'ORDER006'])
    s = book.create_sheet('工單連結')
    s.append(['工單號', '成品料號', '分段式料號'])
    s.append(['LINK001', '98000001-107', '95000001-107'])
    s = book.create_sheet('合併欄位')
    s.append(['機種', 'IC', '料號', '工單號'])
    s.append(['合併機種', 'Elan', 'MERGED', 'MERGED001'])
    s.append([None, None, None, 'MERGED002'])
    for col in 'ABC':
        s.merge_cells(f'{col}2:{col}3')
    # Format in XFD without data: bounds must be based on actual cells.
    s.cell(7, 16384).number_format = '0.00'
    s = book.create_sheet('產品系列')
    s.append(['Model', 'Platform', 'IC', 'Miyabi internal P/N'])
    s.append(['A16', 'Taroko 14', 'Goodix', 'FAMILY-PART'])
    s = book.create_sheet('無對照')
    s.append(['備註', '數量'])
    s.append(['僅保存原始資料', 0])
    s = book.create_sheet('不得讀取')
    s.append(['機種', 'IC', '料號'])
    s.append(['第六頁機種', 'FPC', 'SIXTH-PART'])
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


class LargeWorksheetTests(unittest.TestCase):
    def test_sparse_far_column_and_last_row_are_read_without_rectangle(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '大型對照'
        sheet.append(['機種', 'IC', '料號', '工單'])
        sheet.append(['Model', 'Goodix', 'PART1', 'ORDER1'])
        for column, value in enumerate(['LastModel', 'Elan', 'PART2', 'ORDER2'], 1):
            sheet.cell(100005, column, value)
        sheet.cell(100005, 16384, '保留遠欄')
        stream = io.BytesIO()
        book.save(stream)
        rows = a.excel_sheets(stream.getvalue(), ['大型對照'])[0][1]
        self.assertEqual(len(rows), 100005)
        self.assertEqual(rows[1], ['Model', 'Goodix', 'PART1', 'ORDER1'])
        self.assertEqual(rows[-1][16383], '保留遠欄')
        with tempfile.TemporaryDirectory() as tmp:
            db = a.connect(Path(tmp) / 'large.sqlite3')
            try:
                a.import_excel(db, stream.getvalue(), 'large.xlsx', ['大型對照'])
                self.assertEqual(a.lookup(db, [('order', 'ORDER2')])['候選'][0]['機種'], 'LastModel')
                self.assertEqual(a.stats(db)['索引資料列'], 2)
            finally:
                db.close()


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.db = a.connect(Path(self.tmp.name) / 'test.sqlite3')
        self.payload = fixture()
        self.summary = a.import_excel(self.db, self.payload, 'test.xlsx', FIXTURE_SHEETS)

    def tearDown(self):
        self.db.close()
        self.tmp.cleanup()

    def test_first_five_only_and_idempotence(self):
        self.assertEqual(len(self.summary['sheets']), 5)
        self.assertEqual(a.lookup(self.db, [('part', 'SIXTH-PART')])['狀態'], '查無資料')
        before = self.db.execute('SELECT COUNT(*) FROM records').fetchone()[0]
        self.assertTrue(a.import_excel(self.db, self.payload, 'renamed.xlsx', FIXTURE_SHEETS)['duplicate'])
        self.assertEqual(before, self.db.execute('SELECT COUNT(*) FROM records').fetchone()[0])
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM raw_rows WHERE sheet_id=5').fetchone()[0], 2)

    def test_traditional_simplified_and_aliases(self):
        self.assertEqual(a.role('機型'), 'ic')
        self.assertEqual(a.role('机型'), 'ic')
        self.assertEqual(a.role('工單號碼'), 'order')
        for value in ('繁體測試機種', '繁体测试机种'):
            r = a.lookup(self.db, [('model', value), ('ic', 'ＧＯＯＤＩＸ')])
            self.assertEqual(r['狀態'], '唯一匹配')
            self.assertEqual(r['候選'][0]['IC'], 'Goodix')

    def test_order_link_and_revision_preserved(self):
        r = a.lookup(self.db, [('order', 'link001')])
        self.assertEqual(r['狀態'], '唯一匹配')
        self.assertEqual(r['候選'][0]['機種'], '繁體測試機種')
        self.assertEqual(r['查找依據'][0]['分頁'], '工單連結')
        self.assertEqual(a.lookup(self.db, [('part', '98000001')])['狀態'], '查無資料')

    def test_ambiguous_and_conflicting_conditions(self):
        r = a.lookup(self.db, [('part', 'SHARED')])
        self.assertEqual(r['狀態'], '多組候選')
        self.assertEqual(len(r['候選']), 2)
        r = a.lookup(self.db, [('order', 'ORDER001'), ('ic', 'FPC')])
        self.assertEqual(r['狀態'], '條件不一致')
        self.assertFalse(r['候選'])

    def test_merged_identifiers_and_raw_preserved(self):
        r = a.lookup(self.db, [('order', 'MERGED002')])
        self.assertEqual(r['候選'][0]['機種'], '合併機種')
        self.assertEqual(r['候選'][0]['IC'], 'Elan')
        self.assertEqual(self.db.execute('SELECT cells_json FROM raw_rows WHERE sheet_id=3 AND row_number=3').fetchone()[0],
                         '["", "", "", "MERGED002"]')

    def test_model_platform_separate_and_missing_ic(self):
        r = a.lookup(self.db, [('model', 'A16')])
        self.assertEqual(r['候選'][0]['機種'], 'Taroko 14')
        self.assertEqual(r['候選'][0]['產品系列'], ['A16'])
        self.assertEqual(a.lookup(self.db, [('part', 'UNKNOWN-IC')])['狀態'], '來源未提供完整對應')

    def test_csv_upload_and_excel_export_are_functional(self):
        data = '工单号码,产品料号,机种,机型\nORDER001,,,\n,SHARED,,\n,NO-SUCH-PART,,\n'.encode('utf-8-sig')
        result = a.match_upload(self.db, data, 'upload.csv')
        self.assertEqual([r['狀態'] for r in result['結果']], ['唯一匹配', '多組候選', '查無資料'])
        self.assertEqual(result['結果'][0]['機種'], '繁體測試機種')
        self.assertEqual(result['結果'][1]['IC'], '')
        book = load_workbook(io.BytesIO(a.export_matches(result)), data_only=True)
        self.assertEqual(book.active.cell(2, 4).value, '繁體測試機種')
        self.assertEqual(book.active.max_row, 4)

    def test_xlsx_upload_with_simplified_headers(self):
        book = Workbook()
        book.active.append(['工单号', '料号', '机种', '机型'])
        book.active.append(['ORDER001', '', '', ''])
        buf = io.BytesIO()
        book.save(buf)
        result = a.match_upload(self.db, buf.getvalue(), 'request.xlsx')
        self.assertEqual(result['結果'][0]['IC'], 'Goodix')

    def test_excel_formula_injection_exported_as_text(self):
        result = a.match_upload(self.db, '查询值\n=HYPERLINK("evil")\n'.encode(), 'formula.csv')
        book = load_workbook(io.BytesIO(a.export_matches(result)), data_only=False)
        self.assertFalse(any(c.data_type == 'f' for row in book.active for c in row))

    def test_literal_search_and_no_sql_injection(self):
        self.assertTrue(a.search(self.db, '繁体测试'))
        self.assertFalse(a.search(self.db, "%'; DROP TABLE records;--"))
        self.assertFalse(a.search(self.db, '%'))
        self.assertTrue(a.stats(self.db)['索引資料列'])

    def test_default_selection_is_named_tabs_not_first_five_positions(self):
        book = Workbook()
        book.remove(book.active)
        for name in ['模版', '越南需求', '工单WIP汇总', '0711業務出貨需求']:
            s = book.create_sheet(name)
            s.append(['工單號', 'IC', '機種', '料號'])
            s.append(['C999-999999999', 'WRONG', '不應匯入', 'WRONG-PART'])
        for i, name in enumerate(a.REFERENCE_SHEETS):
            s = book.create_sheet(a.CC.convert(name))
            s.append(['工单号', '机型', '机种', '料号'])
            s.append([f'C515-26010000{i}', 'Goodix', '正確機種', f'9800000{i}-101'])
        payload = io.BytesIO()
        book.save(payload)
        other = a.connect(Path(self.tmp.name) / 'selected.sqlite3')
        result = a.import_excel(other, payload.getvalue(), 'named.xlsx')
        self.assertEqual([a.norm(s['name']) for s in result['sheets']], [a.norm(s) for s in a.REFERENCE_SHEETS])
        self.assertEqual(a.lookup(other, [('part', 'WRONG-PART')])['狀態'], '查無資料')
        self.assertEqual(a.stats(other)['索引資料列'], 5)
        other.close()

    def test_missing_selected_tab_fails_before_writing(self):
        before = self.db.execute('SELECT COUNT(*) FROM sources').fetchone()[0]
        with self.assertRaisesRegex(ValueError, '缺少指定來源分頁'):
            a.import_excel(self.db, self.payload, 'missing.xlsx')
        self.assertEqual(self.db.execute('SELECT COUNT(*) FROM sources').fetchone()[0], before)

    def test_first_work_order_only_sap_prefix_and_placeholders(self):
        self.assertEqual(a.order_numbers('C520-260600014\nSAP:C520-626070046'), ['C520-260600014'])
        self.assertEqual(a.order_numbers('SAP:C520-626070046\nC520-260600014'), ['C520-626070046'])
        self.assertEqual(a.order_numbers('-\nC520-626070046'), [])
        self.assertEqual(a.order_numbers('C530-260600035*50'), ['C530-260600035'])
        self.assertEqual(a.order_numbers('626090006'), ['626090006'])
        for value in ['-', '4月转5月', '板材异常', '52300303-101', 'C531']:
            self.assertEqual(a.order_numbers(value), [])
        book = Workbook()
        book.active.title = '訂單'
        book.active.append(['工單號', '機型', '機種', '料號'])
        book.active.append(['C520-260600014\nSAP:C520-626070046', 'Goodix', '多工單', '98000200-101'])
        payload = io.BytesIO()
        book.save(payload)
        a.import_excel(self.db, payload.getvalue(), 'multi.xlsx', ['訂單'])
        self.assertEqual(a.lookup(self.db, [('order', 'C520-260600014')])['候選'][0]['機種'], '多工單')
        self.assertEqual(a.lookup(self.db, [('order', 'C520-626070046')])['狀態'], '查無資料')

    def test_first_entry_for_model_ic_and_part_keeps_slash_in_name(self):
        book = Workbook()
        book.active.title = '第一筆'
        book.active.append(['工單號', '機型', '機種', '料號', 'DPN'])
        book.active.append(['C515-260100001\nC515-260100002', 'Goodix\nSynaptics',
                            'Slate/Congo 14\n其他機種', '98000398-103\n98000412-103', 'X677X\nPCC61'])
        buf = io.BytesIO()
        book.save(buf)
        a.import_excel(self.db, buf.getvalue(), 'first.xlsx', ['第一筆'])
        result = a.lookup(self.db, [('order', 'C515-260100001')])
        self.assertEqual(result['候選'][0]['IC'], 'Goodix')
        self.assertEqual(result['候選'][0]['機種'], 'Slate/Congo 14')
        self.assertEqual(a.lookup(self.db, [('order', 'C515-260100002')])['狀態'], '查無資料')
        self.assertEqual(a.lookup(self.db, [('part', '98000412-103')])['狀態'], '查無資料')
        self.assertEqual(a.lookup(self.db, [('dpn', 'PCC61')])['狀態'], '查無資料')

    def test_rework_label_resolves_ic_using_exact_part_and_swapped_dpn(self):
        book = Workbook()
        book.active.title = '標準'
        book.active.append(['工單號', 'IC', '機種', '料號'])
        book.active.append(['C515-260100001', 'CY26Goodix', 'CY26-14', '98000398-103'])
        s = book.create_sheet('重工')
        s.append(['工單號', 'IC', '機種', '料號', 'DPN', ''])
        s.append(['C630-626080041', '不良品返修', 'CY26', '95000331-102', '98000398-103', 'C523-260800002'])
        buf = io.BytesIO()
        book.save(buf)
        a.import_excel(self.db, buf.getvalue(), 'rework.xlsx', ['標準', '重工'])
        for order in ['C630-626080041', 'C523-260800002']:
            result = a.lookup(self.db, [('order', order)])
            self.assertEqual(result['候選'][0]['IC'], 'Goodix')
            self.assertEqual(result['候選'][0]['來源'][0]['原始IC／機型'], '不良品返修')
        result = a.lookup(self.db, [('order', 'C515-260100001')])
        self.assertEqual(result['候選'][0]['IC'], 'Goodix')
        self.assertEqual(result['候選'][0]['產品系列'], ['CY26'])

    def test_rework_without_reference_remains_incomplete(self):
        book = Workbook()
        book.active.title = '重工'
        book.active.append(['工單號', 'IC', '機種', '料號'])
        book.active.append(['C630-626080999', '重工测试程式', '無對照機種', '98000999-101'])
        buf = io.BytesIO()
        book.save(buf)
        a.import_excel(self.db, buf.getvalue(), 'unresolved.xlsx', ['重工'])
        result = a.lookup(self.db, [('order', 'C630-626080999')])
        self.assertEqual(result['狀態'], '來源未提供完整對應')
        self.assertEqual(result['候選'][0]['IC'], '')

    def test_rework_conflicting_part_ic_is_not_inferred(self):
        book = Workbook()
        book.active.title = '參考'
        book.active.append(['工單號', 'IC', '機種', '料號'])
        book.active.append(['C515-260100001', 'CY26Goodix', 'CY26-14', '98000412-103'])
        book.active.append(['C515-260100002', 'CY26Synaptics', 'CY26-14', '98000412-103'])
        s = book.create_sheet('重工')
        s.append(['工單號', 'IC', '機種', '料號'])
        s.append(['C630-626080043', '不良品返修', 'CY26', '98000412-103'])
        buf = io.BytesIO()
        book.save(buf)
        a.import_excel(self.db, buf.getvalue(), 'conflict.xlsx', ['參考', '重工'])
        r = a.lookup(self.db, [('order', 'C630-626080043')])
        self.assertEqual(r['狀態'], '來源未提供完整對應')
        self.assertEqual(r['候選'][0]['IC'], '')

    def test_work_order_with_missing_model_is_incomplete_not_absent(self):
        book = Workbook()
        book.active.title = '缺機種'
        book.active.append(['工單號', 'IC', '機種'])
        book.active.append(['C510-260500009', 'Elan', ''])
        buf = io.BytesIO()
        book.save(buf)
        a.import_excel(self.db, buf.getvalue(), 'partial.xlsx', ['缺機種'])
        r = a.lookup(self.db, [('order', 'C510-260500009')])
        self.assertEqual(r['狀態'], '來源未提供完整對應')
        self.assertEqual(r['候選'][0]['IC'], 'Elan')


if __name__ == '__main__':
    unittest.main()
