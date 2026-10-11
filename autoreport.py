"""Import only the first five Excel sheets and resolve model/IC with provenance."""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sqlite3
import unicodedata
import warnings
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from opencc import OpenCC
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries

ROOT = Path(__file__).resolve().parent
DEFAULT_DB = ROOT / 'data' / 'autoreport.sqlite3'
CC = OpenCC('t2s')
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def norm(value):
    """Traditional/simplified, width, case, whitespace, and common dash variants."""
    value = unicodedata.normalize('NFKC', str(value or ''))
    value = value.translate(str.maketrans({'–': '-', '—': '-', '−': '-'}))
    return re.sub(r'\s+', '', CC.convert(value).casefold())


def text(value):
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def available(value):
    return norm(value) not in {'', '-', 'n/a', 'na', 'none', '无', '未知'}


def role(header):
    h = norm(header)
    if h in {'ic', '机型', 'ic型号', 'ic厂商', 'ic品牌'}:
        return 'ic'
    if h in {'机种', 'platform', '产品名称', '产品名', 'model', '机种名称'}:
        return 'model'
    if h in {'dpn', '客户品号', '客户料号', '客户pn', 'customerp/n'}:
        return 'dpn'
    if '工单' in h or h in {'客户订单号', '工令号', 'workorder', 'workorderno'}:
        return 'order'
    if '料号' in h or h in {'产品料号', '产品品号', '品号', 'p/n', 'pn',
                              'miyabiexternalp/n', 'miyabiinternalp/n', 'partnumber'}:
        return 'part'
    if h in {'查询值', '查找值', '关键字', '关键词', 'key', 'query'}:
        return 'any'
    return None


def excel_sheets(payload):
    """Preserve raw cells; expand merged identifiers in a separate working copy.

    Inspect XML bounds so formatting up to XFD does not allocate 16384 columns.
    No formulas, macros, links, or instructions in the file are executed.
    """
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Data Validation extension.*')
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=True,
                             keep_links=False)
    output = []
    with ZipFile(io.BytesIO(payload)) as archive:
        for sheet in list(book)[:5]:
            xml = ET.fromstring(archive.read(sheet._worksheet_path.lstrip('/')))
            cells = [c for c in xml.findall('.//m:c', NS)
                     if any(c.find('m:' + tag, NS) is not None for tag in ('v', 'is', 'f'))]
            max_col = max((column_index_from_string(re.sub('[^A-Za-z]', '', c.attrib['r']))
                           for c in cells), default=1)
            max_row = max((int(re.sub('[^0-9]', '', c.attrib['r'])) for c in cells), default=1)
            if max_col > 512 or max_row > 100000:
                raise ValueError('資料超過支援範圍（每頁 512 欄、100000 列）。')
            raw = [[text(v) for v in r] for r in sheet.iter_rows(
                max_row=max_row, max_col=max_col, values_only=True)]
            expanded = [r.copy() for r in raw]
            for cell in xml.findall('.//m:mergeCell', NS):
                c1, r1, c2, r2 = range_boundaries(cell.attrib['ref'])
                if r1 > max_row or c1 > max_col:
                    continue
                value = raw[r1 - 1][c1 - 1]
                for r in range(r1 - 1, min(r2, max_row)):
                    for c in range(c1 - 1, min(c2, max_col)):
                        if not expanded[r][c]:
                            expanded[r][c] = value
            output.append((sheet.title, raw, expanded))
    book.close()
    return output


def find_header(rows):
    best = None
    for i, row in enumerate(rows[:50]):
        roles = [role(v) for v in row]
        score = sum(r is not None for r in roles)
        if score and (best is None or score > best[0]):
            best = (score, i, row, roles)
    return best


def connect(path=DEFAULT_DB):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(path)
    db.row_factory = sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    db.executescript('''
    CREATE TABLE IF NOT EXISTS sources (
        id INTEGER PRIMARY KEY, filename TEXT NOT NULL, sha256 TEXT UNIQUE NOT NULL,
        imported_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS sheets (
        id INTEGER PRIMARY KEY, source_id INTEGER REFERENCES sources(id),
        position INTEGER NOT NULL, name TEXT NOT NULL, header_row INTEGER,
        raw_rows INTEGER NOT NULL, indexed_rows INTEGER NOT NULL);
    CREATE TABLE IF NOT EXISTS records (
        id INTEGER PRIMARY KEY, sheet_id INTEGER REFERENCES sheets(id), row_number INTEGER,
        model TEXT NOT NULL, ic TEXT NOT NULL, family TEXT NOT NULL,
        parts_json TEXT NOT NULL, dpns_json TEXT NOT NULL, orders_json TEXT NOT NULL,
        raw_json TEXT NOT NULL, search_text TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS lookup_keys (
        kind TEXT NOT NULL, value TEXT NOT NULL, record_id INTEGER REFERENCES records(id),
        PRIMARY KEY(kind, value, record_id));
    CREATE INDEX IF NOT EXISTS lookup_by_record ON lookup_keys(record_id);
    CREATE TABLE IF NOT EXISTS raw_rows (
        sheet_id INTEGER REFERENCES sheets(id), row_number INTEGER,
        cells_json TEXT NOT NULL, PRIMARY KEY(sheet_id, row_number));
    ''')
    return db


def import_excel(db, payload, filename):
    digest = hashlib.sha256(payload).hexdigest()
    previous = db.execute('SELECT id FROM sources WHERE sha256=?', (digest,)).fetchone()
    if previous:
        return {'duplicate': True, 'source_id': previous['id']}
    sheets = excel_sheets(payload)
    summary = []
    with db:
        source_id = db.execute('INSERT INTO sources(filename,sha256,imported_at) VALUES(?,?,?)',
                               (Path(filename).name, digest, datetime.now().isoformat())).lastrowid
        for position, (name, raw, expanded) in enumerate(sheets, 1):
            header = find_header(raw)
            sheet_id = db.execute('INSERT INTO sheets(source_id,position,name,header_row,raw_rows,indexed_rows) '
                                  'VALUES(?,?,?,?,?,0)',
                                  (source_id, position, name, header[1] + 1 if header else None,
                                   sum(any(r) for r in raw))).lastrowid
            db.executemany('INSERT INTO raw_rows VALUES(?,?,?)',
                           [(sheet_id, i, json.dumps(r, ensure_ascii=False))
                            for i, r in enumerate(raw, 1) if any(r)])
            count = 0
            if header:
                _, offset, headers, roles = header
                # English Model is a product family when Platform is also present.
                platforms = [c for c, h in enumerate(headers) if norm(h) == 'platform']
                families = [c for c, h in enumerate(headers) if norm(h) == 'model'] if platforms else []
                for i in range(offset + 1, len(raw)):
                    row = expanded[i]
                    values = {r: [row[c] for c, rr in enumerate(roles) if rr == r and row[c]]
                              for r in ('model', 'ic', 'part', 'dpn', 'order')}
                    if not any(values.values()):
                        continue
                    if any(norm(v) in {'ttl', 'total', '合计', '总计', 'dailycapacity', 'wip天数'}
                           for v in values['model']):
                        continue
                    if any(role(row[c]) == roles[c] for c in range(len(roles))
                           if roles[c] and row[c]):
                        # Skip repeated header rows.
                        if sum(role(row[c]) == roles[c] for c in range(len(roles))
                               if roles[c] and row[c]) >= 2:
                            continue
                    model_values = [row[c] for c in platforms if row[c]] if platforms else values['model']
                    model = model_values[0] if model_values else ''
                    ic = values['ic'][0] if values['ic'] else ''
                    family = next((row[c] for c in families if row[c]), '')
                    # Keep duplicate column names and physical column positions in provenance.
                    original = [{'column': c + 1, 'header': h, 'value': raw[i][c]}
                                for c, h in enumerate(headers) if h or raw[i][c]]
                    record_id = db.execute('INSERT INTO records(sheet_id,row_number,model,ic,family,'
                                           'parts_json,dpns_json,orders_json,raw_json,search_text) '
                                           'VALUES(?,?,?,?,?,?,?,?,?,?)',
                                           (sheet_id, i + 1, model, ic, family,
                                            json.dumps(values['part'], ensure_ascii=False),
                                            json.dumps(values['dpn'], ensure_ascii=False),
                                            json.dumps(values['order'], ensure_ascii=False),
                                            json.dumps(original, ensure_ascii=False), norm(' '.join(row)))).lastrowid
                    keys = set()
                    for kind in ('part', 'dpn', 'order', 'ic'):
                        keys.update((kind, norm(v), record_id) for v in values[kind] if norm(v))
                    keys.update(('model', norm(v), record_id) for v in (model, family) if norm(v))
                    db.executemany('INSERT OR IGNORE INTO lookup_keys VALUES(?,?,?)', sorted(keys))
                    count += 1
            db.execute('UPDATE sheets SET indexed_rows=? WHERE id=?', (count, sheet_id))
            summary.append({'position': position, 'name': name, 'header_row': header[1] + 1 if header else None,
                            'raw_rows': sum(any(r) for r in raw), 'indexed_rows': count})
    return {'duplicate': False, 'source_id': source_id, 'sheets': summary}


RECORD_SQL = '''SELECT r.*, s.name AS sheet, src.filename AS source_file
                FROM records r JOIN sheets s ON s.id=r.sheet_id
                JOIN sources src ON src.id=s.source_id'''


def present(record):
    r = dict(record)
    for field in ('parts', 'dpns', 'orders'):
        r[field] = json.loads(r.pop(field + '_json'))
    r.pop('search_text', None)
    r.pop('raw_json', None)
    return r


def direct_ids(db, kind, value):
    if kind == 'any':
        return {r[0] for r in db.execute('SELECT DISTINCT record_id FROM lookup_keys WHERE value=?', (norm(value),))}
    return {r[0] for r in db.execute('SELECT record_id FROM lookup_keys WHERE kind=? AND value=?', (kind, norm(value)))}


def lookup(db, criteria):
    """Intersect distinct model/IC pairs; never choose the first ambiguous hit.

    Orders without model/IC may resolve through exact part/DPN relations.
    Revision suffixes are preserved; fuzzy text search never auto-fills a result.
    """
    sets, all_candidates, evidence, resolved_criteria = [], {}, {}, []
    for kind, value in criteria:
        if not norm(value):
            continue
        ids = direct_ids(db, kind, value)
        linked = set(ids)
        for record_id in ids:
            rec = db.execute('SELECT model,ic FROM records WHERE id=?', (record_id,)).fetchone()
            if rec['model'] and rec['ic']:
                continue
            # Only identifiers in source rows lacking mappings can bridge to another row.
            for key in db.execute("SELECT kind,value FROM lookup_keys WHERE record_id=? AND kind IN ('part','dpn')",
                                  (record_id,)).fetchall():
                linked.update(direct_ids(db, key['kind'], key['value']))
        pairs = set()
        for record_id in sorted(linked):
            rec = db.execute(RECORD_SQL + ' WHERE r.id=?', (record_id,)).fetchone()
            if not rec['model']:
                continue
            pair = (norm(rec['model']), norm(rec['ic']))
            pairs.add(pair)
            all_candidates.setdefault(pair, {'機種': rec['model'], 'IC': rec['ic'], '產品系列': set(), '來源': {}})
            if rec['family']:
                all_candidates[pair]['產品系列'].add(rec['family'])
            all_candidates[pair]['來源'][record_id] = {
                '檔案': rec['source_file'], '分頁': rec['sheet'], '列號': rec['row_number']}
        # Records with IC should supersede incomplete mappings for the same identifier.
        if any(ic for _, ic in pairs):
            pairs = {p for p in pairs if p[1]}
        sets.append(pairs)
        resolved_criteria.append({'欄位類別': kind, '查找值': str(value), '找到候選組數': len(pairs)})
        for record_id in sorted(ids):
            rec = db.execute(RECORD_SQL + ' WHERE r.id=?', (record_id,)).fetchone()
            evidence[record_id] = {'檔案': rec['source_file'], '分頁': rec['sheet'], '列號': rec['row_number']}
    intersection = set.intersection(*sets) if sets else set()
    candidates = []
    for pair in sorted(intersection):
        c = all_candidates[pair]
        candidates.append({'機種': c['機種'], 'IC': c['IC'], '產品系列': sorted(c['產品系列']),
                           '來源': list(c['來源'].values())})
    if not sets:
        status = '缺少查找欄位'
    elif not intersection:
        status = '條件不一致' if len(sets) > 1 and any(sets) else '查無資料'
    elif len(candidates) > 1:
        status = '多組候選'
    elif not available(candidates[0]['IC']) or not available(candidates[0]['機種']):
        status = '來源未提供完整對應'
    else:
        status = '唯一匹配'
    return {'狀態': status, '候選': candidates, '查找條件': resolved_criteria,
            '查找依據': list(evidence.values())}


def search(db, query, limit=100):
    q = norm(query).replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
    rows = db.execute(RECORD_SQL + " WHERE r.search_text LIKE ? ESCAPE '\\' ORDER BY r.id LIMIT ?",
                      ('%' + q + '%', min(max(int(limit), 1), 500))).fetchall()
    return [present(r) for r in rows]


def stats(db):
    return {'分頁': [dict(r) for r in db.execute('SELECT s.*,src.filename FROM sheets s '
                                                'JOIN sources src ON src.id=s.source_id ORDER BY s.id')],
            '索引資料列': db.execute('SELECT COUNT(*) FROM records').fetchone()[0],
            '機種IC組合': db.execute("SELECT COUNT(*) FROM (SELECT DISTINCT model,ic FROM records WHERE model<>'' AND ic<>'')").fetchone()[0],
            'IC清單': [r[0] for r in db.execute("SELECT DISTINCT ic FROM records WHERE ic<>'' ORDER BY ic")]}


def input_tables(payload, filename):
    if filename.lower().endswith('.xlsx'):
        return excel_sheets(payload)
    if filename.lower().endswith('.csv'):
        for encoding in ('utf-8-sig', 'cp950', 'gb18030'):
            try:
                decoded = payload.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        else:
            raise ValueError('CSV 編碼無法辨識，請使用 UTF-8。')
        rows = list(csv.reader(io.StringIO(decoded)))
        width = max(map(len, rows), default=0)
        if width > 512 or len(rows) > 100000:
            raise ValueError('CSV 超過支援範圍。')
        rows = [r + [''] * (width - len(r)) for r in rows]
        return [(Path(filename).stem, rows, rows)]
    raise ValueError('請上傳 .xlsx 或 .csv 檔案。')


def match_upload(db, payload, filename):
    results, notices = [], []
    for name, raw, expanded in input_tables(payload, filename):
        header = find_header(raw)
        if not header:
            notices.append(f'「{name}」沒有可辨識的查找欄位；請使用工單號、料號、DPN、機種、IC／機型或查詢值。')
            continue
        _, offset, headers, roles = header
        for i in range(offset + 1, len(raw)):
            if not any(raw[i]):
                continue
            criteria = [(r, expanded[i][c]) for c, r in enumerate(roles) if r and expanded[i][c]]
            result = lookup(db, criteria)
            unique = result['候選'][0] if result['狀態'] == '唯一匹配' else {}
            results.append({'分頁': name, '列號': i + 1,
                            '原始資料': [{'欄位': h, '值': raw[i][c]} for c, h in enumerate(headers)],
                            '機種': unique.get('機種', ''), 'IC': unique.get('IC', ''), **result})
    return {'檔名': Path(filename).name, '筆數': len(results), '提示': notices, '結果': results}


def export_matches(result):
    book = Workbook()
    sheet = book.active
    sheet.title = '比對結果'
    sheet.append(['來源分頁', '來源列號', '比對狀態', '機種', 'IC', '候選機種與IC', '查找條件', '資料庫來源', '原始資料'])
    for item in result['結果']:
        sheet.append([item['分頁'], item['列號'], item['狀態'], item['機種'], item['IC'],
                      json.dumps(item['候選'], ensure_ascii=False), json.dumps(item['查找條件'], ensure_ascii=False),
                      json.dumps(item['查找依據'], ensure_ascii=False), json.dumps(item['原始資料'], ensure_ascii=False)])
    # Uploaded content is data; force strings to remain text instead of Excel formulas.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col, width in {'A': 25, 'B': 12, 'C': 18, 'D': 22, 'E': 18, 'F': 55, 'G': 45, 'H': 45, 'I': 55}.items():
        sheet.column_dimensions[col].width = width
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--db', type=Path, default=DEFAULT_DB)
    sub = parser.add_subparsers(dest='command', required=True)
    imp = sub.add_parser('import')
    imp.add_argument('file', type=Path)
    qry = sub.add_parser('query')
    qry.add_argument('value')
    qry.add_argument('--kind', choices=['any', 'order', 'part', 'dpn', 'model', 'ic'], default='any')
    sch = sub.add_parser('search')
    sch.add_argument('value')
    mat = sub.add_parser('match')
    mat.add_argument('file', type=Path)
    mat.add_argument('--out', type=Path, required=True)
    sub.add_parser('stats')
    args = parser.parse_args()
    with connect(args.db) as db:
        if args.command == 'import':
            result = import_excel(db, args.file.read_bytes(), args.file.name)
        elif args.command == 'query':
            result = lookup(db, [(args.kind, args.value)])
        elif args.command == 'search':
            result = search(db, args.value)
        elif args.command == 'match':
            result = match_upload(db, args.file.read_bytes(), args.file.name)
            args.out.parent.mkdir(parents=True, exist_ok=True)
            args.out.write_bytes(export_matches(result))
            result = {k: v for k, v in result.items() if k != '結果'} | {'輸出檔案': str(args.out)}
        else:
            result = stats(db)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
