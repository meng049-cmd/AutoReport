"""Create human-readable outputs and a portable bundle from the imported database."""
import json
from collections import Counter, defaultdict
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import autoreport as a


def style(sheet):
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='157E76')
        cell.font = Font(color='FFFFFF', bold=True)
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = min(60, max(14, max(len(str(c.value or '')) for c in column[:100]) + 2))


def main():
    out = a.ROOT / 'exports'
    out.mkdir(exist_ok=True)
    with a.connect() as db:
        summary = a.stats(db)
        rows = list(db.execute(a.RECORD_SQL + ' ORDER BY r.id'))
        valid = [r for r in rows if a.available(r['model']) and a.available(r['ic'])]
        book = Workbook()
        book.remove(book.active)
        for source in summary['分頁']:
            sheet = book.create_sheet(source['name'][:31])
            for record in db.execute('SELECT row_number,cells_json FROM raw_rows WHERE sheet_id=? ORDER BY row_number', (source['id'],)):
                for col, value in enumerate(json.loads(record['cells_json']), 1):
                    if value:
                        cell = sheet.cell(record['row_number'], col, value)
                        cell.data_type = 's'
            # These tabs preserve original row numbers and values; no invented headers.
            sheet.freeze_panes = f"A{(source['header_row'] or 1) + 1}"
        sheet = book.create_sheet('分頁摘要')
        sheet.append(['順序', '分頁', '表頭列', '非空列數', '可查找列數'])
        for s in summary['分頁']:
            sheet.append([s['position'], s['name'], s['header_row'], s['raw_rows'], s['indexed_rows']])
        style(sheet)
        sheet = book.create_sheet('機種IC對照')
        sheet.append(['機種', 'IC', '產品系列', '料號', 'DPN', '來源分頁', '來源列號', '完整性'])
        mapping = defaultdict(list)
        for r in rows:
            if not r['model'] or not r['ic']:
                continue
            key = (r['model'], r['ic'], r['family'], r['parts_json'], r['dpns_json'], r['sheet'])
            mapping[key].append(str(r['row_number']))
        for key, numbers in sorted(mapping.items()):
            model, ic, family, parts, dpns, source = key
            sheet.append([model, ic, family, '、'.join(json.loads(parts)), '、'.join(json.loads(dpns)), source,
                          '、'.join(numbers), '完整' if a.available(model) and a.available(ic) else '來源未提供完整對應'])
        style(sheet)
        sheet = book.create_sheet('工單與料號索引')
        sheet.append(['來源分頁', '來源列號', '機種', 'IC', '產品系列', '料號', 'DPN', '工單號'])
        for r in rows:
            sheet.append([r['sheet'], r['row_number'], r['model'], r['ic'], r['family'],
                          '、'.join(json.loads(r['parts_json'])), '、'.join(json.loads(r['dpns_json'])),
                          '、'.join(json.loads(r['orders_json']))])
        style(sheet)
        conflicts = []
        for key in db.execute("SELECT DISTINCT value FROM lookup_keys WHERE kind='part' ORDER BY value"):
            result = a.lookup(db, [('part', key[0])])
            if result['狀態'] == '多組候選':
                conflicts.append((key[0], result))
        sheet = book.create_sheet('料號多組候選')
        sheet.append(['料號（正規化）', '候選數', '候選機種／IC', '來源'])
        for value, result in conflicts:
            sheet.append([value, len(result['候選']),
                          '；'.join(c['機種'] + ' / ' + c['IC'] for c in result['候選']),
                          json.dumps([c['來源'] for c in result['候選']], ensure_ascii=False)])
        style(sheet)
        book.save(out / '前五分頁_整理與對照.xlsx')
        template = Workbook()
        sheet = template.active
        sheet.title = '待查詢資料'
        sheet.append(['工單號碼', '產品料號', '機種', 'IC'])
        for row in [['C515-250300004', '', '', ''], ['', '98000320-107', '', ''],
                    ['', 'FP-D01304BE11', '', ''], ['', '找不到的料號', '', ''],
                    ['', '98000239-106', '', ''], ['C515-250300004', '', '', 'Synaptics']]:
            sheet.append(row)
        style(sheet)
        template.save(out / '上傳查詢範例.xlsx')
        payload = (out / '上傳查詢範例.xlsx').read_bytes()
        result = a.match_upload(db, payload, '上傳查詢範例.xlsx')
        (out / '範例_比對結果.xlsx').write_bytes(a.export_matches(result))
        counts = Counter(r['狀態'] for r in result['結果'])
        report = f'''# 工單資料分析與整理

僅讀取附件依原始順序的前五個分頁；第六頁起未匯入或分析。附件文字不作為操作指令。

| 順序 | 分頁 | 原始非空列（含表頭／摘要） | 查找索引列 |
|---|---|---:|---:|
'''
        report += ''.join(f"| {s['position']} | {s['name']} | {s['raw_rows']} | {s['indexed_rows']} |\n" for s in summary['分頁'])
        incomplete = sum(bool(r['model'] and r['ic']) and
                         (not a.available(r['model']) or not a.available(r['ic'])) for r in rows)
        unknown_models = sum(r['model'] == '-' for r in rows)
        unknown_ics = sum(r['ic'] == 'N/A' for r in rows)
        report += f'''
共保存 {sum(s['raw_rows'] for s in summary['分頁']):,} 列原始非空資料，建立 {len(rows):,} 列查找索引。
其中 {len(valid):,} 列提供有效機種與 IC，共 {len({(a.norm(r['model']), a.norm(r['ic'])) for r in valid})} 組正規化後的機種／IC 組合。
找到 {len(conflicts)} 個料號對應多組機種／IC；完整列表在整理檔「料號多組候選」。

## 欄位與比對規則

- 「模版」：有工單／料號資料，原表沒有機種與 IC。以完全相同的料號或 DPN 連結對照資料，不忽略料號版本尾碼。
- 「越南需求」：原始資料全部保留。機種名稱所在欄沒有明確欄名，且未提供 IC，因此不推定對應，不納入自動填值索引。
- 「工单WIP汇总」：機種是 A16、CY24&25、HP 等彙總分類，無 IC，保留原值但不當作完整產品對照。
- 「0711業務出貨需求」：表頭在第 27 列。Platform 作為機種（如 Taroko 14），Model 作為產品系列（如 A16），兩者保留區別；IC 採原始欄位。
- 「工单（A16）」：表頭在第 3 列。使用「機種」「IC」，將不同製程的工單號與 MES 工單號分别建立索引；合併欄位的標識值向合併範圍補齊，但保留原始儲存格資料。
- 「IC」「機型」「机型」依使用者需求視為同一欄位。繁簡轉換用 OpenCC，並統一大小寫、全半形及空白。
- 多個非空查找欄位以候選「機種／IC」組合交集篩選。多組候選、缺資料、條件不一致均不自動填值。
- 文字搜尋只供檢視；部分符合不會自動填入結果。

## 資料品質與限制

來源有 {unknown_models} 列機種為「-」，{unknown_ics} 列 IC 為「N/A」，共 {incomplete} 列已有機種／IC 欄值但不提供完整對應，原值保留。IC 欄另有「FM-3473-007」等型號字串，保留原值，不擅自當成其他 IC 品牌。
第 5 頁第 141 列存在原始公式快取錯誤 `#DIV/0!`，不具有可查找的工單／料號／機種／IC，僅保存於原始資料。
讀取公式快取值，不執行或重算公式；若 Excel 尚未儲存計算結果，空白快取不會被猜測補值。
上傳查詢檔只比對，不更新或覆蓋資料庫。追加來源需使用 README 中的 import 指令；相同檔案雜湊重複匯入不會產生重複資料。

## 驗證

- 工單 `C515-250300004` → Taroko 14 / Goodix。
- 料號 `98000320-107` → Taroko 14 / Goodix。
- 外部料號 `FP-D01304BE11` → Sentry 14\" / Goodix、Thena 14\" / Goodix，標示多組候選。
- 範例檔 6 筆比對結果：{dict(counts)}。
- 已提供資料库、整理 Excel、可上傳查找的本機網頁工具與命令列工具；無需資料庫外掛或外部資料庫服務。
'''
        (out / '分析說明.md').write_text(report, encoding='utf-8')
        print(json.dumps({'valid_rows': len(valid), 'ambiguous_parts': len(conflicts),
                          'sample_results': dict(counts), 'exports': str(out)}, ensure_ascii=False, indent=2))
    with ZipFile(out / 'AutoReport_資料庫與查詢工具.zip', 'w', ZIP_DEFLATED) as bundle:
        for name in ['autoreport.py', 'server.py', 'index.html', 'requirements.txt', 'README.md',
                     'test_autoreport.py', 'build_artifacts.py']:
            bundle.write(a.ROOT / name, 'AutoReport/' + name)
        bundle.write(a.DEFAULT_DB, 'AutoReport/data/autoreport.sqlite3')
        for file in out.iterdir():
            if file.suffix in {'.xlsx', '.md'}:
                bundle.write(file, 'AutoReport/exports/' + file.name)


if __name__ == '__main__':
    main()
