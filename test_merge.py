import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from merge import merge_detail


class MergeTests(unittest.TestCase):
    def test_fail_priority_and_date_process_scope(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'input.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.title = 'Detail'
            sheet.append(['製程', '工單', '料號', '品名規格', '小板序號', '產出時間', '操作人員', '不良項目', '狀態'])
            def row(process, day, status, defect):
                return [process, '001', 'part', 'name', 'serial', datetime(2026, 10, day), 'operator', defect, status]
            sheet.append(row('A', 5, 'GOOD', None))
            sheet.append(row('A', 5, 'FAIL', 'defect'))
            sheet.append(row('A', 5, 'FAIL', 'later'))
            sheet.append(row('A', 6, 'GOOD', None))
            sheet.append(row('B', 5, 'GOOD', None))
            book.save(path)
            result = merge_detail([path], directory)
            self.assertEqual((result['input_rows'], result['output_rows'], result['removed_rows']), (5, 3, 2))
            self.assertEqual(result['filename'], '2026_w41_1005-1006.xlsx')
            output = load_workbook(Path(directory) / result['filename'], read_only=True)
            self.assertEqual(output.sheetnames, ['Detail'])
            rows = list(output['Detail'].iter_rows(min_row=2, values_only=True))
            self.assertEqual(rows[0][7:9], ('defect', 'FAIL'))
            self.assertEqual(rows[-1][5].day, 6)
            output.close()


if __name__ == '__main__':
    unittest.main()
