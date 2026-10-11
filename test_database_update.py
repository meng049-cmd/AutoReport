import hashlib
import importlib.util
import io
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
import autoreport as a

spec = importlib.util.spec_from_file_location('update_database', Path(__file__).parent / 'web/update_database.py')
update = importlib.util.module_from_spec(spec)
spec.loader.exec_module(update)


def workbook(extra=False, missing=False):
    book = Workbook()
    book.remove(book.active)
    for i, name in enumerate(a.REFERENCE_SHEETS[:-1] if missing else a.REFERENCE_SHEETS):
        sheet = book.create_sheet(name)
        sheet.append(['機種', 'IC', '料號', '工單'])
        sheet.append([f'Model{i}', 'Goodix', f'PART{i}', f'ORDER{i}'])
        if extra:
            sheet.append([f'New{i}', 'Elan', f'NEWPART{i}', f'NEWORDER{i}'])
    payload = io.BytesIO()
    book.save(payload)
    return payload.getvalue()


class DatabaseUpdateTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.db = self.root / 'database.sqlite3'
        with a.connect(self.db) as db:
            a.import_excel(db, workbook(), 'original.xlsx')
        self.excel = self.root / 'reference.xlsx'
        self.metadata = self.root / 'update.json'

    def tearDown(self):
        self.temp.cleanup()

    def test_append_preserves_old_records_and_deduplicates_full_workbook(self):
        with a.connect(self.db) as db:
            originals = list(db.execute('SELECT * FROM records ORDER BY id'))
        self.excel.write_bytes(workbook(extra=True))
        result = update.append_excel(self.db, self.excel, self.metadata, 'commit123')
        self.assertEqual(result['added_records'], 5)
        self.assertEqual(result['skipped_records'], 5)
        with a.connect(self.db) as db:
            self.assertEqual([tuple(r) for r in db.execute('SELECT * FROM records WHERE id<=5 ORDER BY id')], [tuple(r) for r in originals])
            self.assertEqual(a.lookup(db, [('order', 'ORDER0')])['狀態'], '唯一匹配')
            self.assertEqual(a.lookup(db, [('order', 'NEWORDER0')])['候選'][0]['機種'], 'New0')
            self.assertEqual(a.stats(db)['索引資料列'], 10)
        metadata_before = self.metadata.read_bytes()
        database_before = self.db.read_bytes()
        self.assertTrue(update.append_excel(self.db, self.excel, self.metadata)['unchanged'])
        self.assertEqual(database_before, self.db.read_bytes())
        self.assertEqual(metadata_before, self.metadata.read_bytes())
        self.assertEqual(json.loads(metadata_before)['source_commit'], 'commit123')

    def test_unchanged_rows_from_a_different_file_do_not_change_database(self):
        self.excel.write_bytes(workbook())
        before = self.db.read_bytes()
        self.assertTrue(update.append_excel(self.db, self.excel, self.metadata)['unchanged'])
        self.assertEqual(before, self.db.read_bytes())
        self.assertFalse(self.metadata.exists())

    def test_invalid_or_missing_sheets_leave_database_intact(self):
        before = hashlib.sha256(self.db.read_bytes()).digest()
        for payload in (b'not an Excel file', workbook(missing=True)):
            self.excel.write_bytes(payload)
            with self.assertRaises(Exception):
                update.append_excel(self.db, self.excel, self.metadata)
            self.assertEqual(before, hashlib.sha256(self.db.read_bytes()).digest())
            self.assertFalse(self.metadata.exists())

    def test_cross_sheet_duplicates_are_skipped_but_conflicting_ic_is_retained(self):
        book = Workbook()
        book.remove(book.active)
        for name in a.REFERENCE_SHEETS:
            sheet = book.create_sheet(name)
            sheet.append(['機種', 'IC', '料號', '工單'])
            sheet.append(['Model0', 'Goodix', 'PART0', 'ORDER0'])
            sheet.append(['Model0', 'Synaptics', 'PART0', 'ORDER0'])
        book.save(self.excel)
        result = update.append_excel(self.db, self.excel, self.metadata)
        self.assertEqual(result['added_records'], 1)
        self.assertEqual(result['skipped_records'], 9)
        with a.connect(self.db) as db:
            actual = a.lookup(db, [('order', 'ORDER0')])
            self.assertEqual(actual['狀態'], '多組候選')
            self.assertEqual({r['IC'] for r in actual['候選']}, {'Goodix', 'Synaptics'})


if __name__ == '__main__':
    unittest.main()
