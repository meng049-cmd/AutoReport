from datetime import datetime, date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell


def merge_detail(paths, output_dir):
    selected = {}
    header = None
    total = 0
    for path in paths:
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = [s for s in book if s.title.strip().lower() == 'detail']
            if len(sheets) != 1:
                raise ValueError(f'{Path(path).name}：找不到唯一的 Detail 分頁。')
            rows = sheets[0].iter_rows(values_only=True)
            current_header = next(rows, None)
            if current_header is None:
                raise ValueError('Detail 分頁沒有標題。')
            if header is None:
                header = current_header
                required = {'製程': 0, '小板序號': 4}
                for title, index in required.items():
                    if len(header) <= index or header[index] != title:
                        raise ValueError('A 欄須為製程，E 欄須為小板序號。')
                if header.count('產出時間') != 1 or header.count('狀態') != 1:
                    raise ValueError('Detail 須包含唯一的「產出時間」及「狀態」欄位。')
                time_index = header.index('產出時間')
                status_index = header.index('狀態')
            elif current_header != header:
                raise ValueError('各檔案 Detail 的欄位標題與順序必須一致。')
            for line, row in enumerate(rows, 2):
                if all(v is None for v in row):
                    continue
                timestamp = row[time_index]
                if isinstance(timestamp, date) and not isinstance(timestamp, datetime):
                    timestamp = datetime.combine(timestamp, datetime.min.time())
                    row = tuple(timestamp if i == time_index else v for i, v in enumerate(row))
                process = str(row[0] or '').strip()
                serial = str(row[4] or '').strip()
                status = str(row[status_index] or '').strip().upper()
                if not isinstance(timestamp, datetime) or not process or not serial or status not in ('GOOD', 'FAIL'):
                    raise ValueError(f'{Path(path).name} 第 {line} 列：日期、製程、小板序號或 GOOD／FAIL 狀態無效。')
                key = (timestamp.date(), process, serial)
                old = selected.get(key)
                if old is None or (status == 'FAIL' and str(old[status_index]).strip().upper() != 'FAIL'):
                    selected[key] = row
                total += 1
        finally:
            book.close()
    if not selected:
        raise ValueError('Detail 沒有可合併的資料。')
    rows = sorted(selected.values(), key=lambda row: row[time_index])
    first, last = rows[0][time_index].date(), rows[-1][time_index].date()
    iso_year, week, _ = first.isocalendar()
    name = f'{iso_year}_w{week}_{first:%m%d}-{last:%m%d}.xlsx'
    output = Path(output_dir) / name
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Detail')
    sheet.freeze_panes = 'A2'
    sheet.append(header)
    for row in rows:
        cells = list(row)
        cell = WriteOnlyCell(sheet, row[time_index])
        cell.number_format = 'yyyy-mm-dd hh:mm:ss'
        cells[time_index] = cell
        sheet.append(cells)
    book.save(output)
    return {'filename': name, 'input_rows': total, 'output_rows': len(rows),
            'removed_rows': total - len(rows), 'start_date': str(first), 'end_date': str(last),
            'fail_rows': sum(str(r[status_index]).strip().upper() == 'FAIL' for r in rows)}
